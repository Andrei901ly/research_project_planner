import streamlit as st
import pandas as pd
import plotly.express as px
from datetime import date, timedelta, datetime
from io import BytesIO

from docx import Document

st.set_page_config(
    page_title="Research Project Planner",
    layout="wide"
)

# -----------------------------
# Session state / temporary data
# -----------------------------
def default_project():
    return {
        "title": "",
        "researcher": "",
        "adviser": "",
        "methodology": "Quantitative",
        "problem": "",
        "objectives": "",
        "questions": "",
        "hypothesis": "",
        "population": "",
        "sample_size": 0,
        "location": "",
        "start_date": date.today(),
        "target_date": date.today() + timedelta(days=90),
    }


def init_state():
    defaults = {
        "project": default_project(),
        "project_form_data": default_project(),
        "tasks": [],
        "literature": [],
        "data_collection": [],
        "chapters": [
            {"chapter": "Chapter 1 - Introduction", "status": "Not Started", "progress": 0},
            {"chapter": "Chapter 2 - Review of Related Literature", "status": "Not Started", "progress": 0},
            {"chapter": "Chapter 3 - Methodology", "status": "Not Started", "progress": 0},
            {"chapter": "Chapter 4 - Results", "status": "Not Started", "progress": 0},
            {"chapter": "Chapter 5 - Conclusion", "status": "Not Started", "progress": 0},
        ],
        "timeline": [
            {"stage": "Topic Selection", "start": date.today(), "deadline": date.today() + timedelta(days=7), "status": "Completed", "progress": 100},
            {"stage": "Literature Review", "start": date.today() + timedelta(days=3), "deadline": date.today() + timedelta(days=21), "status": "Not Started", "progress": 0},
            {"stage": "Research Proposal", "start": date.today() + timedelta(days=14), "deadline": date.today() + timedelta(days=35), "status": "Not Started", "progress": 0},
            {"stage": "Instrument Preparation", "start": date.today() + timedelta(days=30), "deadline": date.today() + timedelta(days=45), "status": "Not Started", "progress": 0},
            {"stage": "Data Collection", "start": date.today() + timedelta(days=42), "deadline": date.today() + timedelta(days=65), "status": "Not Started", "progress": 0},
            {"stage": "Data Analysis", "start": date.today() + timedelta(days=60), "deadline": date.today() + timedelta(days=75), "status": "Not Started", "progress": 0},
            {"stage": "Final Paper", "start": date.today() + timedelta(days=70), "deadline": date.today() + timedelta(days=90), "status": "Not Started", "progress": 0},
            {"stage": "Final Defense", "start": date.today() + timedelta(days=85), "deadline": date.today() + timedelta(days=95), "status": "Not Started", "progress": 0},
        ],
        "notes": "",
        "note_history": [],
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value

init_state()

# -----------------------------
# Helpers
# -----------------------------
def overall_progress():
    values = []
    for t in st.session_state.tasks:
        values.append(100 if t["status"] == "Completed" else t["progress"])
    for c in st.session_state.chapters:
        values.append(c["progress"])
    for x in st.session_state.timeline:
        values.append(x["progress"])
    return round(sum(values) / len(values)) if values else 0

def task_dataframe():
    return pd.DataFrame(st.session_state.tasks)

def excel_download(dataframes):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet, df in dataframes.items():
            df.to_excel(writer, sheet_name=sheet[:31], index=False)
    return output.getvalue()


def docx_download(dataframes):
    document = Document()
    document.add_heading("Research Project Report", level=1)

    for sheet, df in dataframes.items():
        document.add_heading(sheet, level=2)
        if df.empty:
            document.add_paragraph("No data available.")
            continue
        for _, row in df.iterrows():
            items = []
            for key, value in row.items():
                if pd.isna(value):
                    value = ""
                items.append(f"{key}: {value}")
            document.add_paragraph(" | ".join(items))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def normalize_date_value(value):
    if value is None or pd.isna(value):
        return value
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, str):
        try:
            return pd.to_datetime(value).date()
        except Exception:
            return value
    return value


def update_session_records(key, df, date_columns=None):
    if df is None:
        return

    if df.empty:
        st.session_state[key] = []
        return

    records = df.to_dict("records")
    if date_columns:
        for row in records:
            for col in date_columns:
                if col in row:
                    row[col] = normalize_date_value(row[col])
    st.session_state[key] = records

# -----------------------------
# Sidebar
# -----------------------------
st.sidebar.title("Research Planner")

pages = [
    "Dashboard",
    "Research Information",
    "Research Timeline",
    "Tasks",
    "Literature Review",
    "Data Collection",
    "Chapter Progress",
    "Research Notes",
    "Analytics",
    "Export",
]
page = st.sidebar.radio("Navigation", pages)

st.sidebar.divider()
st.sidebar.info(
    "Your information is stored only in the current Streamlit session. "
    "Restarting/reloading the app can clear the data."
)

# -----------------------------
# Dashboard
# -----------------------------
if page == "Dashboard":
    st.title("Research Project Planner")
    st.write("Plan, monitor, and organize your research project in one place.")

    p = st.session_state.project
    progress = overall_progress()
    today = date.today()
    days_left = (p["target_date"] - today).days

    st.subheader("Project Overview")
    if p["title"]:
        st.markdown(f"### {p['title']}")
    else:
        st.warning("No research title has been entered yet.")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Overall Progress", f"{progress}%")
    c2.metric("Tasks", len(st.session_state.tasks))
    c3.metric("Literature Sources", len(st.session_state.literature))
    c4.metric("Days Remaining", max(days_left, 0))

    st.progress(progress / 100)

    st.subheader("Quick Statistics")
    completed_tasks = sum(t["status"] == "Completed" for t in st.session_state.tasks)
    overdue_tasks = sum(
        t["status"] != "Completed" and t["due_date"] < today
        for t in st.session_state.tasks
    )
    completed_stages = sum(x["status"] == "Completed" for x in st.session_state.timeline)

    a, b, c, d = st.columns(4)
    a.metric("Completed Tasks", completed_tasks)
    b.metric("Pending Tasks", len(st.session_state.tasks) - completed_tasks)
    c.metric("Overdue Tasks", overdue_tasks)
    d.metric("Completed Stages", completed_stages)

    st.subheader("Upcoming Deadlines")
    upcoming = []
    for t in st.session_state.tasks:
        if t["status"] != "Completed":
            upcoming.append({
                "Type": "Task",
                "Item": t["task"],
                "Deadline": t["due_date"],
                "Priority": t["priority"]
            })
    for x in st.session_state.timeline:
        if x["status"] != "Completed":
            upcoming.append({
                "Type": "Research Stage",
                "Item": x["stage"],
                "Deadline": x["deadline"],
                "Priority": "High"
            })

    if upcoming:
        upcoming_df = pd.DataFrame(upcoming).sort_values("Deadline").head(8)
        st.dataframe(upcoming_df, use_container_width=True, hide_index=True)
    else:
        st.success("No pending deadlines.")

    st.subheader("Research Workflow")
    st.write(
        "Topic → Literature Review → Proposal → Instrument → "
        "Data Collection → Analysis → Results → Final Paper → Defense"
    )

# -----------------------------
# Research Information
# -----------------------------
elif page == "Research Information":
    st.title("Research Information")
    form_data = st.session_state.get("project_form_data", default_project())

    def save_project_information():
        st.session_state.project = {
            "title": st.session_state["project_title"],
            "researcher": st.session_state["project_researcher"],
            "adviser": st.session_state["project_adviser"],
            "methodology": st.session_state["project_methodology"],
            "start_date": st.session_state["project_start_date"],
            "target_date": st.session_state["project_target_date"],
            "location": st.session_state["project_location"],
            "population": st.session_state["project_population"],
            "sample_size": st.session_state["project_sample_size"],
            "problem": st.session_state["project_problem"],
            "objectives": st.session_state["project_objectives"],
            "questions": st.session_state["project_questions"],
            "hypothesis": st.session_state["project_hypothesis"],
        }
        st.session_state["research_info_saved"] = True

        defaults = default_project()
        st.session_state["project_title"] = defaults["title"]
        st.session_state["project_researcher"] = defaults["researcher"]
        st.session_state["project_adviser"] = defaults["adviser"]
        st.session_state["project_methodology"] = defaults["methodology"]
        st.session_state["project_start_date"] = defaults["start_date"]
        st.session_state["project_target_date"] = defaults["target_date"]
        st.session_state["project_location"] = defaults["location"]
        st.session_state["project_population"] = defaults["population"]
        st.session_state["project_sample_size"] = defaults["sample_size"]
        st.session_state["project_problem"] = defaults["problem"]
        st.session_state["project_objectives"] = defaults["objectives"]
        st.session_state["project_questions"] = defaults["questions"]
        st.session_state["project_hypothesis"] = defaults["hypothesis"]
        st.session_state["project_form_data"] = defaults

    with st.form("research_info_form"):
        c1, c2 = st.columns(2)
        title = c1.text_input("Research Title", key="project_title", value=form_data["title"])
        researcher = c2.text_input("Researcher / Student Name", key="project_researcher", value=form_data["researcher"])

        c1, c2 = st.columns(2)
        adviser = c1.text_input("Research Adviser", key="project_adviser", value=form_data["adviser"])
        methodology_options = ["Quantitative", "Qualitative", "Mixed Methods", "Experimental",
                               "Descriptive", "Correlational", "Other"]
        methodology = c2.selectbox(
            "Research Methodology",
            methodology_options,
            index=methodology_options.index(form_data["methodology"])
            if form_data["methodology"] in methodology_options else 0,
            key="project_methodology"
        )

        c1, c2 = st.columns(2)
        start_date = c1.date_input("Research Start Date", value=form_data["start_date"], key="project_start_date")
        target_date = c2.date_input("Target Completion Date", value=form_data["target_date"], key="project_target_date")

        location = st.text_input("Research Location", key="project_location", value=form_data["location"])
        population = st.text_input("Target Population", key="project_population", value=form_data["population"])
        sample_size = st.number_input("Target Sample Size", min_value=0, value=int(form_data["sample_size"]), key="project_sample_size")

        problem = st.text_area("Statement of the Problem", value=form_data["problem"], key="project_problem")
        objectives = st.text_area("Research Objectives", value=form_data["objectives"], key="project_objectives")
        questions = st.text_area("Research Questions", value=form_data["questions"], key="project_questions")
        hypothesis = st.text_area("Hypothesis", value=form_data["hypothesis"], key="project_hypothesis")

        st.form_submit_button(
            "Save Research Information",
            type="primary",
            on_click=save_project_information,
        )

    if st.session_state.pop("research_info_saved", False):
        st.success("Research information saved temporarily.")

    st.subheader("Research Information Table")
    project_df = pd.DataFrame([st.session_state.project]) if st.session_state.project and any(
        value not in (None, "", 0) for value in st.session_state.project.values()
    ) else pd.DataFrame(columns=list(default_project().keys()))
    edited_project = st.data_editor(
        project_df,
        use_container_width=True,
        hide_index=True,
        key="project_table_editor"
    )
    if not edited_project.empty:
        row = edited_project.iloc[0].to_dict()
        for key in ["start_date", "target_date"]:
            if key in row:
                row[key] = normalize_date_value(row[key])
        st.session_state.project = row

# -----------------------------
# Timeline
# -----------------------------
elif page == "Research Timeline":
    st.title("Research Timeline")

    st.subheader("Add Research Stage")
    with st.form("timeline_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        stage_options = [
            "Topic Selection",
            "Literature Review",
            "Research Proposal",
            "Instrument Preparation",
            "Data Collection",
            "Data Analysis",
            "Final Paper",
            "Final Defense",
            "Other"
        ]
        stage = c1.selectbox("Stage Name", stage_options, key="timeline_stage")
        status = c2.selectbox("Status", ["Not Started", "In Progress", "Completed"], key="timeline_status")

        c1, c2, c3 = st.columns(3)
        start = c1.date_input("Start Date", date.today(), key="timeline_start")
        deadline = c2.date_input("Deadline", date.today() + timedelta(days=7), key="timeline_deadline")
        progress = c3.slider("Progress", 0, 100, 0, key="timeline_progress")

        if st.form_submit_button("Add Stage"):
            st.session_state.timeline.append({
                "stage": stage or "Untitled Stage",
                "start": start,
                "deadline": deadline,
                "status": status,
                "progress": 100 if status == "Completed" else progress,
            })
            st.success("Stage added.")

    st.divider()
    timeline_df = pd.DataFrame(st.session_state.timeline)
    if not timeline_df.empty:
        timeline_df["start"] = pd.to_datetime(timeline_df["start"])
        timeline_df["deadline"] = pd.to_datetime(timeline_df["deadline"])

    edited_timeline = st.data_editor(
        timeline_df,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        key="timeline_table_editor"
    )
    if not edited_timeline.empty:
        update_session_records("timeline", edited_timeline, ["start", "deadline"])
    elif st.session_state.timeline:
        st.session_state.timeline = []

    if not timeline_df.empty:
        fig = px.timeline(
            pd.DataFrame(st.session_state.timeline),
            x_start="start",
            x_end="deadline",
            y="stage",
            color="status",
            hover_data=["progress"]
        )
        fig.update_yaxes(autorange="reversed")
        st.plotly_chart(fig, use_container_width=True)

    if st.button("Clear All Timeline Stages"):
        st.session_state.timeline = []
        st.rerun()

# -----------------------------
# Tasks
# -----------------------------
elif page == "Tasks":
    st.title("Task / To-Do Manager")

    with st.form("task_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        task = c1.text_input("Task", key="task_name")
        category = c2.selectbox(
            "Category",
            ["Planning", "Literature Review", "Proposal", "Data Collection",
             "Data Analysis", "Writing", "Defense", "Other"],
            key="task_category"
        )

        c1, c2, c3 = st.columns(3)
        due_date = c1.date_input("Due Date", date.today() + timedelta(days=7), key="task_due_date")
        priority = c2.selectbox("Priority", ["Low", "Medium", "High"], key="task_priority")
        task_progress = c3.slider("Progress", 0, 100, 0, key="task_progress")

        notes = st.text_input("Task Notes", key="task_notes")

        if st.form_submit_button("Add Task", type="primary"):
            st.session_state.tasks.append({
                "task": task or "Untitled Task",
                "category": category,
                "due_date": due_date,
                "priority": priority,
                "progress": task_progress,
                "status": "Completed" if task_progress == 100 else "Pending",
                "notes": notes,
            })
            st.success("Task added.")

    st.divider()

    task_df = pd.DataFrame(st.session_state.tasks)
    edited_tasks = st.data_editor(
        task_df,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        key="tasks_table_editor"
    )
    if not edited_tasks.empty:
        update_session_records("tasks", edited_tasks, ["due_date"])
    elif st.session_state.tasks:
        st.session_state.tasks = []

    if st.session_state.tasks:
        for i, t in enumerate(st.session_state.tasks):
            with st.container(border=True):
                c1, c2, c3, c4 = st.columns([4, 2, 2, 1])
                c1.write(f"**{t['task']}**")
                c1.caption(f"{t['category']} • {t['priority']} priority • Due {t['due_date']}")
                c2.progress(t["progress"] / 100)
                c2.caption(f"{t['progress']}%")
                new_status = c3.selectbox(
                    "Status",
                    ["Pending", "In Progress", "Completed"],
                    index=["Pending", "In Progress", "Completed"].index(
                        t["status"] if t["status"] in ["Pending", "In Progress", "Completed"] else "Pending"
                    ),
                    key=f"status_{i}"
                )
                if c4.button("Delete", key=f"delete_task_{i}"):
                    st.session_state.tasks.pop(i)
                    st.rerun()

                if new_status != t["status"]:
                    t["status"] = new_status
                    if new_status == "Completed":
                        t["progress"] = 100
                    elif t["progress"] == 100:
                        t["progress"] = 50
                    st.rerun()

    else:
        st.info("No tasks yet. Add your first research task above.")

# -----------------------------
# Literature Review
# -----------------------------
elif page == "Literature Review":
    st.title("Literature Review Tracker")

    with st.form("literature_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        author = c1.text_input("Author(s)", key="literature_author")
        year = c2.number_input("Year", min_value=1900, max_value=2100, value=date.today().year, key="literature_year")

        title = st.text_input("Study / Article Title", key="literature_title")
        source = st.text_input("Journal / Source", key="literature_source")
        topic = st.text_input("Topic / Keywords", key="literature_topic")

        c1, c2 = st.columns(2)
        relevance = c1.selectbox("Relevance", ["Low", "Medium", "High"], key="literature_relevance")
        url = c2.text_input("URL / DOI", key="literature_url")

        findings = st.text_area("Key Findings", key="literature_findings")
        notes = st.text_area("Notes", key="literature_notes")

        if st.form_submit_button("Add Literature Source", type="primary"):
            st.session_state.literature.append({
                "author": author,
                "year": int(year),
                "title": title,
                "source": source,
                "topic": topic,
                "relevance": relevance,
                "url": url,
                "key_findings": findings,
                "notes": notes,
            })
            st.success("Literature source added.")

    st.divider()
    lit_df = pd.DataFrame(st.session_state.literature)
    edited_literature = st.data_editor(
        lit_df,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        key="literature_table_editor"
    )
    if not edited_literature.empty:
        update_session_records("literature", edited_literature)
    elif st.session_state.literature:
        st.session_state.literature = []

    if st.session_state.literature:
        c1, c2 = st.columns(2)
        with c1:
            fig = px.histogram(pd.DataFrame(st.session_state.literature), x="year", title="Studies by Year")
            st.plotly_chart(fig, use_container_width=True)
        with c2:
            fig = px.pie(pd.DataFrame(st.session_state.literature), names="relevance", title="Literature Relevance")
            st.plotly_chart(fig, use_container_width=True)

        if st.button("Clear Literature Review"):
            st.session_state.literature = []
            st.rerun()
    else:
        st.info("Add journal articles, books, theses, or other references above.")

# -----------------------------
# Data Collection
# -----------------------------
elif page == "Data Collection":
    st.title("Data Collection Tracker")

    with st.form("data_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        item = c1.text_input("Data Collection Item", key="data_item")
        method = c2.selectbox("Collection Method", ["Survey", "Interview", "Observation", "Experiment", "Documents", "Other"], key="data_method")

        c1, c2, c3 = st.columns(3)
        target = c1.number_input("Target", min_value=0, value=100, key="data_target")
        collected = c2.number_input("Collected", min_value=0, value=0, key="data_collected")
        status = c3.selectbox("Status", ["Not Started", "In Progress", "Completed"], key="data_status")

        notes = st.text_area("Notes", key="data_notes")

        if st.form_submit_button("Add Data Collection Item"):
            st.session_state.data_collection.append({
                "item": item or "Untitled Item",
                "method": method,
                "target": int(target),
                "collected": int(collected),
                "status": status,
                "notes": notes,
            })
            st.success("Data collection item added.")

    st.divider()

    data_df = pd.DataFrame(st.session_state.data_collection)
    edited_data = st.data_editor(
        data_df,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        key="data_collection_table_editor"
    )
    if not edited_data.empty:
        update_session_records("data_collection", edited_data)
    elif st.session_state.data_collection:
        st.session_state.data_collection = []

    if st.session_state.data_collection:
        data_df = pd.DataFrame(st.session_state.data_collection)
        data_df["progress"] = (
            (data_df["collected"] / data_df["target"].replace(0, 1)) * 100
        ).clip(0, 100).round(1)

        fig = px.bar(
            data_df,
            x="item",
            y=["target", "collected"],
            barmode="group",
            title="Target vs Collected"
        )
        st.plotly_chart(fig, use_container_width=True)

        for i, item_data in enumerate(st.session_state.data_collection):
            if st.button(f"Delete {item_data['item']}", key=f"delete_data_{i}"):
                st.session_state.data_collection.pop(i)
                st.rerun()
    else:
        st.info("Add your survey, interview, observation, or other data collection targets.")

# -----------------------------
# Chapter Progress
# -----------------------------
elif page == "Chapter Progress":
    st.title("Chapter Progress Tracker")

    for i, chapter in enumerate(st.session_state.chapters):
        st.markdown(f"### {chapter['chapter']}")
        c1, c2 = st.columns([2, 5])
        new_status = c1.selectbox(
            "Status",
            ["Not Started", "In Progress", "Completed"],
            index=["Not Started", "In Progress", "Completed"].index(chapter["status"]),
            key=f"chapter_status_{i}"
        )
        new_progress = c2.slider(
            "Progress",
            0, 100, int(chapter["progress"]),
            key=f"chapter_progress_{i}"
        )

        if new_status == "Completed":
            new_progress = 100

        chapter["status"] = new_status
        chapter["progress"] = new_progress
        st.progress(new_progress / 100)

    chapter_df = pd.DataFrame(st.session_state.chapters)
    edited_chapters = st.data_editor(
        chapter_df,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        key="chapters_table_editor"
    )
    if not edited_chapters.empty:
        update_session_records("chapters", edited_chapters)

    st.success(f"Overall chapter progress: {round(sum(c['progress'] for c in st.session_state.chapters) / len(st.session_state.chapters))}%")

# -----------------------------
# Notes
# -----------------------------
elif page == "Research Notes":
    st.title("Research Notes")
    st.caption("Use this space for research ideas, adviser comments, meeting notes, and reminders.")

    def save_research_note():
        note = st.session_state.get("research_notes_input", "").strip()
        if note:
            st.session_state.note_history.append({
                "Note": note,
                "Date": date.today(),
                "Time": datetime.now().strftime("%I:%M:%S %p"),
            })
            st.session_state.notes = note
            st.session_state["research_notes_input"] = ""

    notes_value = st.session_state.get("research_notes_input", "")
    current_notes = st.text_area(
        "Notes",
        value=notes_value,
        height=400,
        placeholder="Write your research notes here...",
        key="research_notes_input"
    )
    st.session_state.notes = current_notes

    st.button("Save Notes", on_click=save_research_note)

    if st.session_state.note_history:
        st.subheader("Saved Notes")
        notes_table = pd.DataFrame(
            st.session_state.note_history,
            columns=["Note", "Date", "Time"],
        )
        if "notes_table_version" not in st.session_state:
            st.session_state.notes_table_version = 0
        selected_note = st.dataframe(
            notes_table,
            use_container_width=True,
            hide_index=True,
            on_select="rerun",
            selection_mode="single-row",
            key=f"saved_notes_table_{st.session_state.notes_table_version}",
        )

        def reset_note_table_selection():
            st.session_state.notes_table_version += 1

        @st.dialog("Research Note", on_dismiss=reset_note_table_selection)
        def show_note_details(note_record):
            st.write(note_record["Note"])
            st.divider()
            st.write(f"**Date:** {note_record['Date']}")
            st.write(f"**Time:** {note_record['Time']}")

        selected_rows = selected_note.selection.rows
        if selected_rows:
            show_note_details(
                st.session_state.note_history[selected_rows[0]]
            )
    else:
        st.info("Save a note to add it to the notes table.")

# -----------------------------
# Analytics
# -----------------------------
elif page == "Analytics":
    st.title("Research Analytics")

    progress = overall_progress()
    st.metric("Overall Research Progress", f"{progress}%")
    st.progress(progress / 100)

    # Task analytics
    st.subheader("Task Analytics")
    if st.session_state.tasks:
        task_df = pd.DataFrame(st.session_state.tasks)

        c1, c2 = st.columns(2)
        with c1:
            fig = px.pie(task_df, names="status", title="Tasks by Status")
            st.plotly_chart(fig, use_container_width=True)
        with c2:
            fig = px.bar(task_df, x="category", color="status", title="Tasks by Category")
            st.plotly_chart(fig, use_container_width=True)
    else:
        st.info("Add tasks to display task analytics.")

    # Chapter analytics
    st.subheader("Chapter Analytics")
    chapter_df = pd.DataFrame(st.session_state.chapters)
    fig = px.bar(
        chapter_df,
        x="chapter",
        y="progress",
        title="Chapter Completion",
        range_y=[0, 100]
    )
    st.plotly_chart(fig, use_container_width=True)

    # Timeline analytics
    st.subheader("Research Stage Analytics")
    timeline_df = pd.DataFrame(st.session_state.timeline)
    if not timeline_df.empty:
        fig = px.bar(
            timeline_df,
            x="stage",
            y="progress",
            color="status",
            title="Research Stage Progress",
            range_y=[0, 100]
        )
        st.plotly_chart(fig, use_container_width=True)

    # Data collection analytics
    st.subheader("Data Collection Analytics")
    if st.session_state.data_collection:
        dc_df = pd.DataFrame(st.session_state.data_collection)
        dc_df["progress"] = (
            dc_df["collected"] / dc_df["target"].replace(0, 1) * 100
        ).clip(0, 100).round(1)
        fig = px.bar(
            dc_df,
            x="item",
            y="progress",
            title="Data Collection Progress",
            range_y=[0, 100]
        )
        st.plotly_chart(fig, use_container_width=True)
    else:
        st.info("Add data collection targets to display this chart.")

# -----------------------------
# Export
# -----------------------------
elif page == "Export":
    st.title("Export Research Project")

    st.write("Download your temporary project data as an Excel workbook.")

    project_df = pd.DataFrame([st.session_state.project])
    timeline_df = pd.DataFrame(st.session_state.timeline)
    tasks_df = pd.DataFrame(st.session_state.tasks)
    literature_df = pd.DataFrame(st.session_state.literature)
    data_df = pd.DataFrame(st.session_state.data_collection)
    chapters_df = pd.DataFrame(st.session_state.chapters)
    notes_df = pd.DataFrame(
        st.session_state.note_history,
        columns=["Note", "Date", "Time"],
    )

    workbook = excel_download({
        "Project": project_df,
        "Timeline": timeline_df,
        "Tasks": tasks_df,
        "Literature": literature_df,
        "Data Collection": data_df,
        "Chapters": chapters_df,
        "Notes": notes_df,
    })

    st.download_button(
        "Download Research Project Excel",
        data=workbook,
        file_name="research_project_plan.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )

    docx_file = docx_download({
        "Project": project_df,
        "Timeline": timeline_df,
        "Tasks": tasks_df,
        "Literature": literature_df,
        "Data Collection": data_df,
        "Chapters": chapters_df,
        "Notes": notes_df,
    })
    st.download_button(
        "Download Research Project as DOCX",
        data=docx_file,
        file_name="research_project_report.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="secondary"
    )

    st.subheader("Export Preview")
    st.write(f"**Research Title:** {st.session_state.project['title'] or 'Not set'}")
    st.write(f"**Overall Progress:** {overall_progress()}%")
    st.write(f"**Tasks:** {len(st.session_state.tasks)}")
    st.write(f"**Literature Sources:** {len(st.session_state.literature)}")
    st.write(f"**Data Collection Items:** {len(st.session_state.data_collection)}")

    st.divider()
    st.warning(
        "This application does not use a database. Data exists only in the current "
        "Streamlit session, so download the Excel file if you want to keep a copy."
    )

# -----------------------------
# Footer
# -----------------------------
st.sidebar.divider()